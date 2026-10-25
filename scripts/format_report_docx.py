#!/usr/bin/env python3
"""Apply the final table and image-caption formatting to the Dấu Xưa report."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Inches
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_border(cell, color: str = "808080", size: str = "6") -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)

    for edge_name in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def shade_cell(cell, fill: str = "E7E6E6") -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def disable_repeating_header(row) -> None:
    """Keep the first row styled as a header without repeating it on later pages."""
    row_properties = row._tr.get_or_add_trPr()
    for header_property in list(row_properties.findall(qn("w:tblHeader"))):
        row_properties.remove(header_property)


def keep_row_on_one_page(row) -> None:
    """Move a whole row to the next page instead of splitting its text."""
    row_properties = row._tr.get_or_add_trPr()
    cannot_split = row_properties.find(qn("w:cantSplit"))
    if cannot_split is None:
        cannot_split = OxmlElement("w:cantSplit")
        row_properties.append(cannot_split)
    cannot_split.set(qn("w:val"), "1")


def format_tables(document: Document) -> None:
    for table in document.tables:
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            disable_repeating_header(row)
            keep_row_on_one_page(row)
            for cell in row.cells:
                set_cell_border(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index == 0:
                    shade_cell(cell)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
                        for run in paragraph.runs:
                            run.bold = True


def format_image_captions(document: Document) -> int:
    """Center figure captions and keep them in regular, non-italic text."""
    caption_style = document.styles["Image Caption"]
    caption_style.font.italic = False
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    caption_count = 0
    for paragraph in document.paragraphs:
        if paragraph._element.xpath('.//w:drawing'):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.keep_together = True
        if paragraph.text.startswith("Hình "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_together = True
            for run in paragraph.runs:
                run.italic = False
            caption_count += 1
    return caption_count


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()

    document = Document(args.docx)
    for section in document.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
    format_tables(document)
    caption_count = format_image_captions(document)
    document.save(args.docx)
    print(
        f"Formatted {len(document.tables)} tables and {caption_count} image captions "
        f"in {args.docx}"
    )


if __name__ == "__main__":
    main()
